"""
qc_engine.py - reference loading, image OCR, field extraction, QC decisions,
result Excel writer.  No GUI code in here.
"""
import csv
import os
import re
import threading
import unicodedata
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from difflib import SequenceMatcher

import cv2
import numpy as np
import pandas as pd

import config as cfg

FIELDS = ["id", "name", "taluk", "a2021", "a2024", "a2019"]
FIELD_LABEL = {"id": "Wetland ID", "name": "Wetland Name", "taluk": "Taluk",
               "a2021": "2021 Area", "a2024": "2024 Area", "a2019": "2019 Area"}
EVIDENCE_SUFFIX = {"id": "ID", "name": "Name", "taluk": "Taluk",
                   "a2021": "2021", "a2024": "2024", "a2019": "2019"}
HEADER_FIELDS = ("id", "name", "taluk")
LEGEND_FIELDS = ("a2021", "a2024", "a2019")
YEAR_FIELD = {"2021": "a2021", "2024": "a2024", "2019": "a2019"}

YES, NO, MANUAL = "YES", "NO", "MANUAL CHECK"
PASS, MISMATCH, MANUAL_CHECK = "PASS", "MISMATCH", "MANUAL CHECK"

YEAR_RE = re.compile(r"(?<!\d)(2019|2021|2024)(?!\d)")
NUM_RE = re.compile(r"(?<![0-9.])(\d[\d,]*)(?:\s?\.\s?(\d+))?(?![0-9])")
ID_TOKEN_RE = re.compile(cfg.ID_TOKEN_PATTERN)
STOP_RE = re.compile(cfg.STOP_LABEL, re.I)
AREA_KW_RE = re.compile(cfg.AREA_KEYWORDS, re.I)
NAME_RES = [re.compile(p, re.I) for p in cfg.NAME_LABELS]
TALUK_RES = [re.compile(p, re.I) for p in cfg.TALUK_LABELS]
ANY_LABEL_RE = re.compile("|".join(cfg.NAME_LABELS + cfg.TALUK_LABELS +
                                   [cfg.STOP_LABEL, r"wetland\s*(?:id|code)"]), re.I)


# ============================================================================
# Text / number helpers
# ============================================================================
def norm_text(s):
    if s is None:
        return ""
    s = unicodedata.normalize("NFKC", str(s)).lower()
    s = re.sub(r"[\-_/\\.,;:()\[\]{}'\"`\u2019\u2013\u2014|]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def norm_key(s):
    return re.sub(r"[^a-z0-9]", "", str(s).lower())


def fold_id(s):
    """Upper-case, drop spaces, and map O/I/L -> 0/1/1 in the digit part (OCR noise)."""
    s = re.sub(r"\s+", "", str(s)).upper()
    m = re.match(r"([A-Z]*)(.*)$", s)
    return m.group(1) + m.group(2).translate(str.maketrans("OIL", "011"))


def cell_value(v):
    if v is None:
        return None
    if isinstance(v, float) and np.isnan(v):
        return None
    if isinstance(v, str) and not v.strip():
        return None
    return v


def display(v):
    v = cell_value(v)
    if v is None:
        return ""
    if isinstance(v, float):
        try:
            return format(Decimal(str(v)), "f")
        except InvalidOperation:
            return str(v)
    return str(v).strip()


def parse_excel_number(v):
    """-> (Decimal|None, reason). reason is '' when OK."""
    v = cell_value(v)
    if v is None:
        return None, "blank"
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        try:
            return Decimal(str(v)), ""
        except InvalidOperation:
            return None, "not numeric"
    s = str(v).strip().lower()
    s = re.sub(r"(sq\.?\s*km|hectares?|\bha\b)", "", s)
    s = re.sub(r"\s+", "", s)
    if re.fullmatch(r"\d+,\d{1,2}", s):
        return None, "ambiguous decimal comma"
    s = s.replace(",", "")
    if re.fullmatch(r"-?\d+(\.\d+)?", s):
        return Decimal(s), ""
    return None, "not numeric"


def clean_id(v):
    v = cell_value(v)
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", "", str(v)).upper()


# ============================================================================
# Reference (Excel / CSV) loading
# ============================================================================
class RefError(Exception):
    pass


@dataclass
class RefTable:
    records: dict = field(default_factory=dict)   # WETLAND_ID -> [ {field: raw} ]
    colmap: dict = field(default_factory=dict)    # field -> header text
    missing: list = field(default_factory=list)   # fields with no column anywhere
    n_rows: int = 0

    def summary(self):
        found = ", ".join(f"{FIELD_LABEL[f]}='{h}'" for f, h in self.colmap.items())
        miss = ("  | NOT FOUND: " + ", ".join(FIELD_LABEL[f] for f in self.missing)) if self.missing else ""
        return f"Excel: {self.n_rows} rows | {found}{miss}"


def _read_frames(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                with open(path, newline="", encoding=enc) as f:
                    sample = f.read(4096)
                    f.seek(0)
                    try:
                        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                    except csv.Error:
                        dialect = csv.excel
                    rows = [r for r in csv.reader(f, dialect)]
                break
            except UnicodeDecodeError:
                continue
        else:
            raise RefError("Cannot decode CSV file.")
        width = max((len(r) for r in rows), default=0)
        return {"csv": [r + [None] * (width - len(r)) for r in rows]}
    if ext in (".xlsx", ".xlsm", ".xls"):
        try:
            frames = pd.read_excel(path, sheet_name=None, header=None, dtype=object)
        except ImportError as e:
            raise RefError(f"Missing library to read {ext}: {e}")
        return {name: df.values.tolist() for name, df in frames.items()}
    raise RefError("Unsupported reference file type (use .xlsx, .xls or .csv).")


def detect_columns(headers):
    """headers: list of header strings -> {field: column_index}"""
    norm = [norm_key(h) for h in headers]
    found, used = {}, set()
    for f in FIELDS:                      # pass 1: exact aliases
        for alias in cfg.COLUMN_ALIASES[f]:
            a = norm_key(alias)
            idx = next((i for i, h in enumerate(norm) if h == a and i not in used), None)
            if idx is not None:
                found[f] = idx
                used.add(idx)
                break
    for f in FIELDS:                      # pass 2: loose patterns
        if f in found:
            continue
        rx = re.compile(cfg.COLUMN_LOOSE[f])
        idx = next((i for i, h in enumerate(norm) if h and i not in used and rx.search(h)), None)
        if idx is not None:
            found[f] = idx
            used.add(idx)
    return found


def load_reference(path):
    frames = _read_frames(path)
    table = RefTable()
    for sheet, rows in frames.items():
        best = None
        for r in range(min(len(rows), cfg.HEADER_SCAN_ROWS)):
            headers = [display(c) for c in rows[r]]
            cols = detect_columns(headers)
            if "id" in cols and (best is None or len(cols) > len(best[1])):
                best = (r, cols, headers)
                if len(cols) == len(FIELDS):
                    break
        if best is None:
            continue
        hr, cols, headers = best
        for f, i in cols.items():
            table.colmap.setdefault(f, headers[i])
        for row in rows[hr + 1:]:
            wid = clean_id(row[cols["id"]]) if cols["id"] < len(row) else ""
            if not wid:
                continue
            rec = {f: (row[cols[f]] if f in cols and cols[f] < len(row) else None)
                   for f in FIELDS if f != "id"}
            rec["_cols"] = set(cols)
            rec["_sheet"] = sheet
            table.records.setdefault(wid, []).append(rec)
            table.n_rows += 1
    if not table.records:
        raise RefError("No table with a Wetland ID column was found in the file. "
                       "Check the header names (see COLUMN_ALIASES in config.py).")
    table.missing = [f for f in FIELDS if f not in table.colmap]
    return table


def resolve_id(stem, records):
    """-> (wetland_id | None, found_in_reference)"""
    s = re.sub(r"\s+", "", stem).upper()
    if s in records:
        return s, True
    hits = [i for i in records if len(i) >= 4 and i in s]
    if hits:
        return max(hits, key=len), True
    m = re.search(r"[A-Z]{2}\d{4,9}", s)
    if m:
        return m.group(0), False
    return None, False


# ============================================================================
# Image helpers / OCR
# ============================================================================
def imread_unicode(path):
    try:
        data = np.fromfile(path, dtype=np.uint8)
        return cv2.imdecode(data, cv2.IMREAD_COLOR)
    except Exception:
        return None


def imwrite_unicode(path, img, quality=90):
    ok, buf = cv2.imencode(".jpg", img, [cv2.IMWRITE_JPEG_QUALITY, quality])
    if ok:
        buf.tofile(path)


_tls = threading.local()


def get_engine():
    if not hasattr(_tls, "engine"):
        try:
            # Modern RapidOCR package; works with current Python versions.
            from rapidocr import RapidOCR
        except ImportError:
            # Backward compatibility for older environments.
            from rapidocr_onnxruntime import RapidOCR
        _tls.engine = RapidOCR()
    return _tls.engine


def run_ocr(engine, img):
    out = engine(img)
    if isinstance(out, tuple):                       # rapidocr_onnxruntime
        return [(b, t, float(s)) for b, t, s in (out[0] or [])]
    if getattr(out, "txts", None) is None:           # newer rapidocr
        return []
    return [(b, t, float(s)) for b, t, s in zip(out.boxes, out.txts, out.scores)]


def preprocess(crop, variant):
    h, w = crop.shape[:2]
    s = max(0.5, min(3.0, cfg.OCR_TARGET_WIDTH / float(w)))
    img = cv2.resize(crop, None, fx=s, fy=s,
                     interpolation=cv2.INTER_CUBIC if s > 1 else cv2.INTER_AREA)
    if variant == "base":
        return img, s
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
    if variant == "clahe":
        g = cv2.createCLAHE(clipLimit=2.5, tileGridSize=(8, 8)).apply(gray)
        g = cv2.medianBlur(g, 3)
    elif variant == "otsu":
        g = cv2.GaussianBlur(gray, (3, 3), 0)
        _, g = cv2.threshold(g, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
    else:  # "sharp"
        g = cv2.addWeighted(gray, 1.8, cv2.GaussianBlur(gray, (0, 0), 3), -0.8, 0)
        g = cv2.adaptiveThreshold(g, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
                                  cv2.THRESH_BINARY, 31, 15)
    return cv2.cvtColor(g, cv2.COLOR_GRAY2BGR), s


@dataclass
class Item:
    x0: float
    y0: float
    x1: float
    y1: float
    text: str
    conf: float

    @property
    def cy(self):
        return (self.y0 + self.y1) / 2

    @property
    def h(self):
        return max(1.0, self.y1 - self.y0)


class Row:
    def __init__(self, items):
        self.items = sorted(items, key=lambda i: i.x0)
        text, spans = "", []
        for it in self.items:
            if text:
                text += " "
            s = len(text)
            text += it.text.strip()
            spans.append((s, len(text), it))
        self.text, self.spans = text, spans
        self.cy = sum(i.cy for i in self.items) / len(self.items)
        self.h = sum(i.h for i in self.items) / len(self.items)
        self.mean_conf = sum(i.conf for i in self.items) / len(self.items)
        self.bbox = (min(i.x0 for i in self.items), min(i.y0 for i in self.items),
                     max(i.x1 for i in self.items), max(i.y1 for i in self.items))

    def conf_at(self, off):
        for s, e, it in self.spans:
            if s <= off < e:
                return it.conf
        return self.mean_conf

    def conf_from(self, off):
        cs = [it.conf for s, e, it in self.spans if e > off]
        return sum(cs) / len(cs) if cs else self.mean_conf


def build_rows(items):
    groups = []
    for it in sorted(items, key=lambda i: i.cy):
        if groups:
            g = groups[-1]
            gcy = sum(x.cy for x in g) / len(g)
            gh = sum(x.h for x in g) / len(g)
            if abs(it.cy - gcy) <= 0.5 * max(it.h, gh):
                g.append(it)
                continue
        groups.append([it])
    return [Row(g) for g in groups if g]


def get_regions(img):
    h, w = img.shape[:2]
    name = cfg.ACTIVE_LAYOUT
    if name == "auto" or name not in cfg.LAYOUTS:
        name = "landscape" if w >= h else "portrait"
    return cfg.LAYOUTS[name]


def region_box(img, frac):
    h, w = img.shape[:2]
    x0, y0 = int(max(0, frac[0]) * w), int(max(0, frac[1]) * h)
    x1, y1 = int(min(1, frac[2]) * w), int(min(1, frac[3]) * h)
    return x0, y0, max(x0 + 1, x1), max(y0 + 1, y1)


def region_rows(engine, img, frac, variant, log=None, tag=""):
    """OCR one cropped region; returned coordinates are in ORIGINAL image pixels."""
    x0, y0, x1, y1 = region_box(img, frac)
    crop = img[y0:y1, x0:x1]
    if crop.size == 0:
        return []
    proc, s = preprocess(crop, variant)
    items = []
    for box, text, conf in run_ocr(engine, proc):
        if not str(text).strip():
            continue
        xs = [p[0] for p in box]
        ys = [p[1] for p in box]
        items.append(Item(x0 + min(xs) / s, y0 + min(ys) / s,
                          x0 + max(xs) / s, y0 + max(ys) / s, str(text), conf))
    rows = build_rows(items)
    if log is not None:
        log.append(f"  [{tag}/{variant}] " + " | ".join(r.text for r in rows))
    return rows


# ============================================================================
# Field extraction from OCR rows
# ============================================================================
@dataclass
class Reading:
    value: str
    conf: float
    bbox: tuple = None
    num: Decimal = None
    decimals: int = 0
    ambiguous: bool = False
    note: str = ""


def _union(b1, b2):
    return (min(b1[0], b2[0]), min(b1[1], b2[1]), max(b1[2], b2[2]), max(b1[3], b2[3]))


def extract_id(rows, wid=None):
    cands = []
    for row in rows:
        for m in ID_TOKEN_RE.finditer(row.text):
            cands.append(Reading(fold_id(m.group(0)), row.conf_at(m.start()), row.bbox))
    if not cands:
        return None
    if wid:
        for c in cands:
            if c.value == fold_id(wid):
                return c
    return cands[0]


def _clean_value(v):
    sm = STOP_RE.search(v)
    if sm:
        v = v[:sm.start()]
    return v.strip(" :-\u2013\u2014=|")


def extract_label(rows, label_res):
    for lre in label_res:
        for i, row in enumerate(rows):
            m = lre.search(row.text)
            if not m:
                continue
            tail = row.text[m.end():]
            stripped = tail.lstrip(" :-\u2013\u2014=|.")
            off = m.end() + len(tail) - len(stripped)
            val = _clean_value(stripped)
            if val:
                return Reading(val, row.conf_from(off), row.bbox)
            nxt = rows[i + 1] if i + 1 < len(rows) else None    # value on the next line
            if nxt and nxt.cy - row.cy <= 2.5 * row.h and not ANY_LABEL_RE.search(nxt.text):
                val = _clean_value(nxt.text)
                if val:
                    return Reading(val, nxt.mean_conf, _union(row.bbox, nxt.bbox))
    return None


def _reading_from_parts(ip, fp, conf, bbox):
    ip = str(ip).replace(",", "")
    try:
        num = Decimal(ip + ("." + fp if fp else ""))
    except InvalidOperation:
        return None
    return Reading(ip + ("." + fp if fp else ""), conf, bbox, num=num,
                   decimals=len(fp) if fp else 0)


def _number_candidates(row, start, end):
    """Return all numeric candidates in this year's own text segment."""
    seg = row.text[start:end]
    out = []
    for m in NUM_RE.finditer(seg):
        ip, fp = m.group(1), m.group(2)

        if fp is None and re.fullmatch(r"(19|20)\d\d", ip):
            continue

        conf = row.conf_at(start + m.start())

        if fp is None and re.fullmatch(r"\d+,\d{1,2}", ip):
            out.append(Reading(ip, conf, row.bbox, ambiguous=True,
                               note="decimal comma/thousand separator ambiguous"))
            continue

        r = _reading_from_parts(ip, fp, conf, row.bbox)
        if r:
            out.append(r)
    return out


def _fractional_continuation(text):
    """Detect a next-line fractional fragment such as '.85 Ha' or ',85 Ha'."""
    s = str(text).strip()
    m = re.match(r"^[\.,]\s*(\d{1,6})(?:\s*(?:ha|hectare|hectares|\(ha\)))?\b", s, re.I)
    return m.group(1) if m else None


def _find_number(row, start, end):
    cands = _number_candidates(row, start, end)
    if not cands:
        return None

    good = [c for c in cands if not c.ambiguous]
    if not good:
        return cands[0]

    # Prefer a complete decimal if OCR also produced a partial integer token.
    decimals = [c for c in good if c.decimals > 0]
    if decimals:
        return max(decimals, key=lambda c: c.conf)

    return max(good, key=lambda c: c.conf)


def extract_areas(rows):
    """
    Extract 2019 / 2021 / 2024 independently.

    Fixes split-decimal OCR such as:
      2019 - 1.85 Ha
      2019 - 1 . 85 Ha
      2019 - 1
             .85 Ha
    """
    cands = {y: [] for y in YEAR_FIELD}

    for i, row in enumerate(rows):
        ms = list(YEAR_RE.finditer(row.text))
        if not ms:
            continue

        # 'Ha' may itself be moved to the next OCR row.
        nearby_has_area_kw = bool(AREA_KW_RE.search(row.text))
        if not nearby_has_area_kw and i + 1 < len(rows):
            nearby_has_area_kw = bool(AREA_KW_RE.search(rows[i + 1].text))
        if not nearby_has_area_kw:
            continue

        for k, m in enumerate(ms):
            year = m.group(1)
            end = ms[k + 1].start() if k + 1 < len(ms) else len(row.text)

            seg_cands = _number_candidates(row, m.end(), end)
            good = [c for c in seg_cands if not c.ambiguous]

            r = None
            if good:
                decimal_good = [c for c in good if c.decimals > 0]
                r = max(decimal_good, key=lambda c: c.conf) if decimal_good \
                    else max(good, key=lambda c: c.conf)
            elif seg_cands:
                r = seg_cands[0]

            # Join integer + fractional fragment across adjacent OCR rows.
            if r is not None and not r.ambiguous and r.decimals == 0 and i + 1 < len(rows):
                nxt = rows[i + 1]
                if not YEAR_RE.search(nxt.text) and nxt.cy - row.cy <= 2.8 * row.h:
                    frac = _fractional_continuation(nxt.text)
                    if frac:
                        merged = _reading_from_parts(
                            str(r.value).replace(",", ""),
                            frac,
                            min(r.conf, nxt.mean_conf),
                            _union(row.bbox, nxt.bbox)
                        )
                        if merged:
                            merged.note = "joined split decimal OCR fragments"
                            r = merged

            # If the value is wholly on the next OCR row.
            if r is None and k == len(ms) - 1 and i + 1 < len(rows):
                nxt = rows[i + 1]
                if not YEAR_RE.search(nxt.text) and nxt.cy - row.cy <= 2.8 * row.h:
                    r = _find_number(nxt, 0, len(nxt.text))
                    if r:
                        r.bbox = _union(row.bbox, nxt.bbox)

            if r:
                cands[year].append(r)

    out = {}
    for y, cs in cands.items():
        if not cs:
            out[y] = None
            continue

        good = [c for c in cs if not c.ambiguous]
        distinct = {c.num for c in good}

        if len(distinct) > 1:
            # Prefer a value repeatedly confirmed by multiple preprocessing variants.
            groups = {}
            for c in good:
                groups.setdefault(c.num, []).append(c)
            ranked = sorted(groups.values(),
                            key=lambda g: (len(g), max(x.conf for x in g)),
                            reverse=True)
            if ranked and len(ranked[0]) >= 2 and (
                len(ranked) == 1 or len(ranked[0]) > len(ranked[1])
            ):
                out[y] = max(ranked[0], key=lambda c: c.conf)
            else:
                out[y] = Reading(
                    " / ".join(c.value for c in good),
                    min(c.conf for c in good),
                    cs[0].bbox,
                    ambiguous=True,
                    note=f"conflicting {y} values on map"
                )
        elif good:
            out[y] = max(good, key=lambda c: c.conf)
        else:
            out[y] = cs[0]

    return out


# ============================================================================
# Decision logic
# ============================================================================
def decide(match, rkey, near, readings, final):
    """-> (status, representative_reading, retry_helpful)"""
    valid = [r for r in readings if r is not None and r.value and not r.ambiguous]
    if not valid:
        return MANUAL, None, True
    eq = [r for r in valid if match(r)]
    ne = [r for r in valid if not match(r)]
    if eq:
        best = max(eq, key=lambda r: r.conf)
        if len(eq) > len(ne) and best.conf >= cfg.MATCH_MIN_CONF:
            return YES, best, False
        return MANUAL, best, True
    groups = {}
    for r in ne:
        groups.setdefault(rkey(r), []).append(r)
    top = max(groups.values(), key=lambda g: (len(g), max(x.conf for x in g)))
    rep = max(top, key=lambda r: r.conf)
    if len(groups) > 1 and len(top) < 2:
        return MANUAL, rep, True                                  # readings disagree
    if near and near(rep) and rep.conf < cfg.HIGH_CONF:
        return MANUAL, rep, True                                  # probably OCR noise
    if len(top) >= 2 and rep.conf >= cfg.MIN_NO_CONF:
        return NO, rep, False                                     # confirmed by 2 variants
    if final and len(ne) == 1 and rep.conf >= cfg.HIGH_CONF:
        return NO, rep, False
    return MANUAL, rep, True


def evaluate(f, rec, rl, wid, final):
    """-> (status, map_display_text, retry_helpful, note)"""
    fallback = next((r for r in rl if r is not None and r.value), None)
    notes = "; ".join(sorted({r.note for r in rl if r is not None and r.note}))

    if f == "id":
        target = fold_id(wid)
        res = decide(lambda r: fold_id(r.value) == target, lambda r: fold_id(r.value),
                     None, rl, final)
        if res[1] is None and not fallback and not cfg.ID_REQUIRED_IN_MAP:
            return YES, "", False, "ID taken from filename (not printed in map)"
        status, rep, retry = res
        return status, (rep or fallback).value if (rep or fallback) else "", retry, \
            ("" if status == YES else (notes or ("Wetland ID not readable in map" if not (rep or fallback) else "")))

    if f in ("name", "taluk"):
        if not rec["_cols"].issuperset({f}):
            return MANUAL, (fallback.value if fallback else ""), False, f"{FIELD_LABEL[f]} column not found in Excel"
        ex = norm_text(cell_value(rec[f]))
        if not ex:
            return MANUAL, (fallback.value if fallback else ""), False, f"Excel {FIELD_LABEL[f]} is blank"
        near = lambda r: SequenceMatcher(None, norm_text(r.value), ex).ratio() >= cfg.FUZZY_RATIO
        status, rep, retry = decide(lambda r: norm_text(r.value) == ex,
                                    lambda r: norm_text(r.value), near, rl, final)
        shown = (rep or fallback)
        return status, shown.value if shown else "", retry, \
            ("" if status == YES else (notes or ("" if shown else f"{FIELD_LABEL[f]} not found/readable on map")))

    # area fields
    if f not in rec["_cols"]:
        return MANUAL, (fallback.value if fallback else ""), False, f"{FIELD_LABEL[f]} column not found in Excel"
    ex, why = parse_excel_number(rec[f])
    if ex is None:
        return MANUAL, (fallback.value if fallback else ""), False, f"Excel {FIELD_LABEL[f]} {why}"

    def match(r):
        e = ex
        if cfg.ROUND_EXCEL_TO_MAP_PRECISION:
            ed = -ex.as_tuple().exponent
            if ed > r.decimals:
                e = ex.quantize(Decimal(1).scaleb(-r.decimals), rounding=ROUND_HALF_UP)
        return e == r.num

    status, rep, retry = decide(match, lambda r: r.num, None, rl, final)
    shown = (rep or fallback)
    return status, shown.value if shown else "", retry, \
        ("" if status == YES else (notes or ("" if shown else f"{FIELD_LABEL[f]} not found/readable on map")))


# ============================================================================
# Per-map processing
# ============================================================================
@dataclass
class QCResult:
    filename: str
    path: str
    wid: str = ""
    checks: dict = field(default_factory=lambda: {f: MANUAL for f in FIELDS})
    excel: dict = field(default_factory=lambda: {f: "" for f in FIELDS})
    mapv: dict = field(default_factory=lambda: {f: "" for f in FIELDS})
    overall: str = MANUAL_CHECK
    notes: str = ""
    raw_ocr: str = ""

    def finalize(self):
        bad = [f for f in FIELDS if self.checks[f] == NO]
        man = [f for f in FIELDS if self.checks[f] == MANUAL]
        self.overall = MISMATCH if bad else (MANUAL_CHECK if man else PASS)
        return self

    @property
    def mismatch_fields(self):
        return ", ".join(FIELD_LABEL[f] for f in FIELDS if self.checks[f] == NO)


def _excel_display(res, rec, wid):
    res.excel["id"] = wid
    for f in FIELDS[1:]:
        res.excel[f] = display(rec.get(f)) if rec else ""


def _save_evidence(img, res, bboxes, regions, review_dir):
    os.makedirs(review_dir, exist_ok=True)
    h, w = img.shape[:2]
    for f in FIELDS:
        bb = bboxes.get(f)
        if bb is None:
            bb = region_box(img, regions["header" if f in HEADER_FIELDS else "legend"])
        pad = 14
        x0, y0 = max(0, int(bb[0]) - pad), max(0, int(bb[1]) - pad)
        x1, y1 = min(w, int(bb[2]) + pad), min(h, int(bb[3]) + pad)
        if x1 > x0 and y1 > y0:
            imwrite_unicode(os.path.join(review_dir, f"{res.wid}_{EVIDENCE_SUFFIX[f]}.jpg"),
                            img[y0:y1, x0:x1])


def process_map(path, ref, review_dir):
    fn = os.path.basename(path)
    res = QCResult(filename=fn, path=path)
    stem = os.path.splitext(fn)[0]
    wid, in_ref = resolve_id(stem, ref.records)
    res.wid = wid or stem
    if not wid:
        res.notes = "No Wetland ID found in filename"
        return res.finalize()
    if not in_ref:
        res.checks["id"] = NO
        res.mapv["id"] = wid
        res.notes = "Wetland ID not found in Excel"
        return res.finalize()

    recs = ref.records[wid]
    rec = recs[0]
    _excel_display(res, rec, wid)
    if len(recs) > 1:
        sig = lambda r: tuple(norm_text(display(r.get(f))) for f in FIELDS[1:])
        if len({sig(r) for r in recs}) > 1:
            res.notes = f"Duplicate Wetland ID in Excel ({len(recs)} rows) with different values"
            return res.finalize()

    img = imread_unicode(path)
    if img is None or img.size == 0:
        res.notes = "Unreadable image"
        return res.finalize()

    regions = get_regions(img)
    engine = get_engine()
    log = [f"{fn}  (ID {wid})"]
    readings = {f: [] for f in FIELDS}
    bboxes = {}
    variants = cfg.VARIANTS[:max(1, cfg.MAX_VARIANTS)]
    need_h = need_l = True
    outcome = {}
    try:
        for vi, var in enumerate(variants):
            final = vi == len(variants) - 1
            if need_h:
                rows = region_rows(engine, img, regions["header"], var, log, "header")
                readings["id"].append(extract_id(rows, wid))
                readings["name"].append(extract_label(rows, NAME_RES))
                readings["taluk"].append(extract_label(rows, TALUK_RES))
            if need_l:
                rows = region_rows(engine, img, regions["legend"], var, log, "legend")
                areas = extract_areas(rows)
                for y, f in YEAR_FIELD.items():
                    readings[f].append(areas[y])
            outcome = {f: evaluate(f, rec, readings[f], wid, final) for f in FIELDS}
            need_h = any(outcome[f][0] == MANUAL and outcome[f][2] for f in HEADER_FIELDS)
            need_l = any(outcome[f][0] == MANUAL and outcome[f][2] for f in LEGEND_FIELDS)
            if not (need_h or need_l):
                break
    except Exception as e:
        res.notes = f"OCR failure: {e}"
        res.raw_ocr = "\n".join(log)
        return res.finalize()

    notes = []
    for f in FIELDS:
        status, shown, _, note = outcome[f]
        res.checks[f] = status
        res.mapv[f] = shown if shown else ("" if status == YES else "(not read)")
        if note:
            notes.append(f"{FIELD_LABEL[f]}: {note}")
        last = next((r for r in reversed(readings[f]) if r is not None and r.bbox), None)
        if last:
            bboxes[f] = last.bbox
    res.notes = " | ".join(notes)
    res.raw_ocr = "\n".join(log)
    res.finalize()
    if res.overall != PASS:
        try:
            _save_evidence(img, res, bboxes, regions, review_dir)
        except Exception as e:
            res.notes += f" | evidence not saved: {e}"
    return res


def list_maps(folder):
    files = [os.path.join(folder, n) for n in os.listdir(folder)
             if n.lower().endswith(tuple(cfg.IMAGE_EXTS))]
    return sorted(files, key=lambda p: [int(t) if t.isdigit() else t.lower()
                                        for t in re.split(r"(\d+)", os.path.basename(p))])


def run_batch(ref, maps, out_dir, stop_evt, on_start=None, on_result=None):
    review_dir = os.path.join(out_dir, cfg.REVIEW_DIR_NAME)

    def job(p):
        if stop_evt.is_set():
            return None
        if on_start:
            on_start(os.path.splitext(os.path.basename(p))[0])
        try:
            return process_map(p, ref, review_dir)
        except Exception as e:                      # one bad image must never stop the batch
            r = QCResult(filename=os.path.basename(p), path=p,
                         wid=os.path.splitext(os.path.basename(p))[0])
            r.notes = f"Error: {e}"
            return r.finalize()

    results = []
    with ThreadPoolExecutor(max_workers=max(1, cfg.WORKERS)) as ex:
        futs = [ex.submit(job, p) for p in maps]
        for fut in as_completed(futs):
            r = fut.result()
            if r is None:
                continue
            results.append(r)
            if on_result:
                on_result(r)
    results.sort(key=lambda r: [int(t) if t.isdigit() else t.lower()
                                for t in re.split(r"(\d+)", r.filename)])
    return results


# ============================================================================
# Output Excel
# ============================================================================
def write_output(results, out_dir):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    fills = {"YES": "C6EFCE", "PASS": "C6EFCE", "NO": "FFC7CE", "MISMATCH": "FFC7CE",
             "MANUAL CHECK": "FFEB9C"}
    fills = {k: PatternFill("solid", start_color=v, end_color=v) for k, v in fills.items()}

    wb = Workbook()
    ws = wb.active
    ws.title = "QC Result"
    headers = ["Wetland ID", "Wetland ID Check", "Excel Wetland Name", "Map Wetland Name",
               "Wetland Name Check", "Excel Taluk / Tahasil", "Map Taluk / Tahasil", "Taluk Check",
               "Excel 2021 Area", "Map 2021 Area", "2021 Area Check",
               "Excel 2024 Area", "Map 2024 Area", "2024 Area Check",
               "Excel 2019 Area", "Map 2019 Area", "2019 Area Check",
               "Overall Result", "Mismatch Fields", "Map Filename", "Notes"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color="305496", end_color="305496")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    check_cols = []
    for r in results:
        row = [r.wid, r.checks["id"]]
        for f in FIELDS[1:]:
            row += [r.excel[f], r.mapv[f], r.checks[f]]
        row += [r.overall, r.mismatch_fields, r.filename, r.notes]
        ws.append(row)
    check_cols = [2, 5, 8, 11, 14, 17, 18]
    for row in ws.iter_rows(min_row=2):
        for ci in check_cols:
            cell = row[ci - 1]
            if cell.value in fills:
                cell.fill = fills[cell.value]
                cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.columns, 1):
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)

    sm = wb.create_sheet("Summary")
    cnt = lambda s: sum(1 for r in results if r.overall == s)
    rows = [("Total Maps Checked", len(results)), ("Total PASS", cnt(PASS)),
            ("Total MISMATCH", cnt(MISMATCH)), ("Total MANUAL CHECK", cnt(MANUAL_CHECK)), ("", ""),
            ("Wetland ID Mismatch", sum(r.checks["id"] == NO for r in results)),
            ("Wetland Name Mismatch", sum(r.checks["name"] == NO for r in results)),
            ("Taluk Mismatch", sum(r.checks["taluk"] == NO for r in results)),
            ("2021 Area Mismatch", sum(r.checks["a2021"] == NO for r in results)),
            ("2024 Area Mismatch", sum(r.checks["a2024"] == NO for r in results)),
            ("2019 Area Mismatch", sum(r.checks["a2019"] == NO for r in results))]
    for a, b in rows:
        sm.append([a, b])
    for c in sm["A"]:
        c.font = Font(bold=True)
    sm.column_dimensions["A"].width = 28
    sm.column_dimensions["B"].width = 12

    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, cfg.OUTPUT_EXCEL_NAME)
    try:
        wb.save(path)
    except PermissionError:          # file is open in Excel
        import time
        path = os.path.join(out_dir, time.strftime("Wetland_QC_Result_%Y%m%d_%H%M%S.xlsx"))
        wb.save(path)
    try:                              # raw OCR text for debugging (kept out of the GUI)
        with open(os.path.join(out_dir, cfg.DEBUG_LOG_NAME), "w", encoding="utf-8") as f:
            f.write("\n\n".join(r.raw_ocr for r in results if r.raw_ocr))
    except Exception:
        pass
    return path


# ============================================================================
# Calibration helper (used by the GUI "Calibrate" button)
# ============================================================================
def debug_map(path):
    img = imread_unicode(path)
    if img is None:
        raise RuntimeError("Cannot read image.")
    regions = get_regions(img)
    engine = get_engine()
    report, ann = [], img.copy()
    thick = max(2, img.shape[1] // 400)
    for name, colour in (("header", (0, 0, 255)), ("legend", (255, 0, 0))):
        x0, y0, x1, y1 = region_box(img, regions[name])
        cv2.rectangle(ann, (x0, y0), (x1, y1), colour, thick)
        cv2.putText(ann, name, (x0 + 8, y0 + 12 + thick * 10), cv2.FONT_HERSHEY_SIMPLEX,
                    thick * 0.6, colour, thick)
    hrows = region_rows(engine, img, regions["header"], "base")
    lrows = region_rows(engine, img, regions["legend"], "base")
    report.append("=== HEADER region OCR rows ===")
    report += [f"  ({r.mean_conf:.2f}) {r.text}" for r in hrows]
    report.append("=== LEGEND region OCR rows ===")
    report += [f"  ({r.mean_conf:.2f}) {r.text}" for r in lrows]
    idr, nm, tk = extract_id(hrows), extract_label(hrows, NAME_RES), extract_label(hrows, TALUK_RES)
    areas = extract_areas(lrows)
    fmt = lambda r: f"{r.value}  (conf {r.conf:.2f}){'  AMBIGUOUS' if r.ambiguous else ''}" if r else "-- not found --"
    report += ["=== EXTRACTED ===", f"  Wetland ID : {fmt(idr)}", f"  Name       : {fmt(nm)}",
               f"  Taluk      : {fmt(tk)}"]
    report += [f"  {y} Area  : {fmt(areas[y])}" for y in ("2021", "2024", "2019")]
    return ann, "\n".join(report)
